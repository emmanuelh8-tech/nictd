# build-cost-word.py — NICTD Data Mining Phase cost schedule + Data Team (Word).
import os, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..")
OUT_DIR = os.path.join(BASE, "deliverables")
rows = json.load(open(os.path.join(OUT_DIR, "_mining_computed.json")))["rows"]

NAVY = RGBColor(0x12, 0x26, 0x3A); TEAL = RGBColor(0x1E, 0x8A, 0x8A)
GREY = RGBColor(0x5C, 0x6B, 0x78); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX, TEAL_HEX, CREAM_HEX, GOLD_HEX = "12263A", "1E8A8A", "F7F4EE", "C8102E"
CTR = WD_ALIGN_PARAGRAPH.CENTER

# ---- computed figures (mirror of the Excel; verified) ----
rate = {g: 50 + (g - 1) * 50 for g in range(1, 7)}
fdays = dict(enum=20, kii=20, sup=25, lead=30)
field_cost = sum(r["enumerators"]*fdays["enum"]*rate[1] + r["kii"]*fdays["kii"]*rate[2]
                 + r["supervisors"]*fdays["sup"]*rate[3] + r["research_lead"]*fdays["lead"]*rate[4] for r in rows)
data_team = [
    ("Data Team Lead", "Owns the processing pipeline, review queue & final sign-off", 5, 1, 45),
    ("Sampling Statistician", "Weighting, county estimates, precision & error checks", 5, 1, 25),
    ("Database / CAPI Engineer", "Builds CAPI instrument, sync server, upload pipeline to NICTD", 4, 2, 30),
    ("GIS / Mapping Analyst", "Enumeration-area frames, geocoding, choropleth prep", 4, 1, 20),
    ("Data Quality Analyst", "Validation rules, cleaning, outlier & consistency review", 3, 4, 25),
    ("Data Processing Clerk", "Coding open responses, data entry backup, reconciliation", 2, 6, 15),
]
dt_cost = sum(hc*days*rate[g] for _, _, g, hc, days in data_team)
dt_hc = sum(hc for _, _, g, hc, days in data_team)
mgmt = 1*55*rate[6]
labour = field_cost + dt_cost + mgmt
cont = round(labour*0.10)
grand = labour + cont
tot_enum = sum(r["enumerators"] for r in rows); tot_kii = sum(r["kii"] for r in rows)
tot_sup = sum(r["supervisors"] for r in rows); tot_lead = sum(r["research_lead"] for r in rows)

doc = Document()
doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)
def rn(p, text, size=10, bold=False, color=None, italic=False, white=False):
    r = p.add_run(text); r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if white: r.font.color.rgb = WHITE
    elif color: r.font.color.rgb = color
    return r
def heading(text, size=14, color=NAVY, sb=12):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(4)
    rn(p, text, size=size, bold=True, color=color); return p
def para(text, size=10, color=None, bold=False, italic=False):
    p = doc.add_paragraph(); rn(p, text, size=size, color=color, bold=bold, italic=italic); return p
def set_cell(cell, text, bold=False, color=None, size=9, align=None, white=False):
    cell.text = ""; p = cell.paragraphs[0]
    if align: p.alignment = align
    rn(p, str(text), size=size, bold=bold, color=color, white=white)
def money(x): return f"${x:,.0f}"
def table(cols, widths, fill=NAVY_HEX):
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(cols):
        set_cell(t.rows[0].cells[j], c, bold=True, white=True, size=9); shade(t.rows[0].cells[j], fill)
    t._w = widths; return t
def add_row(t, cells, aligns=None, zebra=False, bold=False, fill=None):
    row = t.add_row().cells
    for j, v in enumerate(cells):
        set_cell(row[j], v, align=(aligns[j] if aligns else None), size=9, bold=bold, white=bool(fill))
        if fill: shade(row[j], fill)
        elif zebra: shade(row[j], CREAM_HEX)
    return row
def fixw(t):
    for row in t.rows:
        for j, w in enumerate(t._w): row.cells[j].width = Inches(w)

# ===== cover =====
p = doc.add_paragraph(); rn(p, "National ICT Database of Liberia (NICTD)", size=20, bold=True, color=NAVY)
p = doc.add_paragraph(); rn(p, "Data Mining Phase — Cost Schedule & Data Team", size=13, color=TEAL)
para("Labour costed on the client pay ladder: base US$50/day at the lowest grade, rising US$50 per grade. Companion live workbook: NICTD-Cost-Schedule.xlsx (editable rates, days & headcount).", size=9.5, color=GREY, italic=True)

# ===== pay ladder =====
heading("1. Pay grade ladder")
t = table(["Grade", "Daily rate", "Roles at this grade"], [1.0, 1.1, 4.4])
ladder = [
    ("G1", "$50", "Enumerator (field data collector)"),
    ("G2", "$100", "Facility / KII collector · Data processing clerk"),
    ("G3", "$150", "Field supervisor · Data quality analyst"),
    ("G4", "$200", "County research lead · Database/CAPI engineer · GIS analyst"),
    ("G5", "$250", "Sampling statistician · Data Team lead"),
    ("G6", "$300", "Survey / Project Director"),
]
for i, (g, rt, roles) in enumerate(ladder):
    add_row(t, [g, rt, roles], aligns=[CTR, CTR, None], zebra=(i % 2 == 1))
fixw(t)

# ===== field team cost =====
heading("2. Field team labour (by county)")
para("Cost = headcount x paid days x grade rate. Paid days: enumerators & KII 20, supervisors 25, research leads 30.", size=9.5, color=GREY, italic=True)
t = table(["County", "Enum.", "KII", "Superv.", "Leads", "County cost"], [1.5, 0.75, 0.6, 0.85, 0.7, 1.3], fill=TEAL_HEX)
for i, r in enumerate(rows):
    cost = r["enumerators"]*fdays["enum"]*rate[1] + r["kii"]*fdays["kii"]*rate[2] + r["supervisors"]*fdays["sup"]*rate[3] + r["research_lead"]*fdays["lead"]*rate[4]
    add_row(t, [r["county"], r["enumerators"], r["kii"], r["supervisors"], r["research_lead"], money(cost)],
            aligns=[None, CTR, CTR, CTR, CTR, CTR], zebra=(i % 2 == 1))
add_row(t, ["FIELD TEAM TOTAL", tot_enum, tot_kii, tot_sup, tot_lead, money(field_cost)],
        aligns=[None, CTR, CTR, CTR, CTR, CTR], fill=NAVY_HEX)
fixw(t)

doc.add_page_break()

# ===== data team =====
heading("3. Data Team — post-collection processing (NEW)")
para("A dedicated central team that turns raw field data into published NICTD indicators: cleaning, validation, weighting, coding, geocoding, database engineering and the bulk upload into the platform's review queue.", size=10)
t = table(["Role", "Responsibility", "Grade", "No.", "Days", "Cost"], [1.7, 3.0, 0.6, 0.5, 0.5, 1.0], fill=TEAL_HEX)
for i, (name, resp, g, hc, days) in enumerate(data_team):
    add_row(t, [name, resp, f"G{g}", hc, days, money(hc*days*rate[g])],
            aligns=[None, None, CTR, CTR, CTR, CTR], zebra=(i % 2 == 1))
add_row(t, ["DATA TEAM TOTAL", "", "", dt_hc, "", money(dt_cost)], aligns=[None, None, CTR, CTR, CTR, CTR], fill=NAVY_HEX)
fixw(t)
para("Workflow: field CAPI data -> nightly sync -> Data Quality Analysts clean & de-duplicate -> Statistician weights and computes county estimates -> GIS analyst maps -> Database/CAPI engineer formats to the platform CSV (indicator_code, county, year, value, source) -> Data Team Lead uploads via the Content Studio, passing automated validation and the human review queue -> live on NICTD.", size=9.5, color=GREY, italic=True)

# ===== summary =====
heading("4. Cost summary")
t = table(["Cost component", "Basis", "USD"], [2.5, 2.9, 1.3])
for label, basis, val, bold, fill in [
    ("Field team labour", "118 enumerators, 33 KII, 31 supervisors, 15 leads", money(field_cost), False, None),
    ("Data team labour", f"{dt_hc}-person central processing team", money(dt_cost), False, None),
    ("Central management", "Survey/Project Director, G6, 55 days", money(mgmt), False, None),
    ("Labour subtotal", "field + data team + management", money(labour), True, TEAL_HEX),
    ("Contingency (10%)", "editable in the workbook", money(cont), False, None),
    ("Non-labour (optional)", "transport, tablets, per-diem, training — fill if applicable", "$0", False, None),
    ("GRAND TOTAL", "labour subtotal + contingency + non-labour", money(grand), True, NAVY_HEX),
]:
    add_row(t, [label, basis, val], aligns=[None, None, CTR], bold=bold, fill=fill)
fixw(t)

heading("5. Workforce recap", sb=10)
t = table(["Group", "Headcount"], [3.5, 1.6])
for label, n in [("Field data collectors (enumerators)", tot_enum), ("Facility / KII collectors", tot_kii),
                 ("Field supervisors", tot_sup), ("County research leads", tot_lead),
                 ("Data team (central)", dt_hc), ("Project director", 1),
                 ("TOTAL WORKFORCE", tot_enum+tot_kii+tot_sup+tot_lead+dt_hc+1)]:
    add_row(t, [label, n], aligns=[None, CTR], bold=(label.startswith("TOTAL")), fill=(NAVY_HEX if label.startswith("TOTAL") else None))
fixw(t)

para("")
para("Assumptions (all editable in NICTD-Cost-Schedule.xlsx): base rate $50 and $50 step; paid days per role as above; 10% contingency; non-labour set to $0 by default. Change the base rate, step, days or headcount in the blue cells and every total recalculates.", size=9, color=GREY, italic=True)

path = os.path.abspath(os.path.join(OUT_DIR, "NICTD-Cost-Schedule.docx"))
doc.save(path)
print("WROTE", path)
print(f"grand=${grand:,.0f} labour=${labour:,.0f} field=${field_cost:,.0f} datateam=${dt_cost:,.0f} dt_hc={dt_hc}")
