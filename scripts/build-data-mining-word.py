# build-data-mining-word.py — NICTD Data Mining Phase brief + full field questionnaire (Word).
import os, json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.join(os.path.dirname(__file__), "..")
OUT_DIR = os.path.join(BASE, "deliverables")
os.makedirs(OUT_DIR, exist_ok=True)
data = json.load(open(os.path.join(OUT_DIR, "_mining_computed.json")))
rows, totals = data["rows"], data["totals"]
researchers_total = data["researchers_total"]

NAVY = RGBColor(0x12, 0x26, 0x3A); TEAL = RGBColor(0x1E, 0x8A, 0x8A)
GREY = RGBColor(0x5C, 0x6B, 0x78); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX, TEAL_HEX, CREAM_HEX, GOLD_HEX = "12263A", "1E8A8A", "F7F4EE", "C8102E"

doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(10)

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)

def run(p, text, size=10, bold=False, color=None, italic=False, white=False):
    r = p.add_run(text); r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if white: r.font.color.rgb = WHITE
    elif color: r.font.color.rgb = color
    return r

def heading(text, size=14, color=NAVY, space_before=10):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(4)
    run(p, text, size=size, bold=True, color=color); return p

def para(text, size=10, color=None, bold=False, italic=False):
    p = doc.add_paragraph(); run(p, text, size=size, color=color, bold=bold, italic=italic); return p

def set_cell(cell, text, bold=False, color=None, size=9, align=None, white=False, italic=False):
    cell.text = ""; p = cell.paragraphs[0]
    if align: p.alignment = align
    run(p, str(text), size=size, bold=bold, color=color, white=white, italic=italic)

def table(cols, widths, header_fill=NAVY_HEX):
    t = doc.add_table(rows=1, cols=len(cols)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(cols):
        set_cell(t.rows[0].cells[j], c, bold=True, white=True, size=9); shade(t.rows[0].cells[j], header_fill)
    t._widths = widths
    return t

def add_row(t, cells, zebra=False, aligns=None):
    row = t.add_row().cells
    for j, v in enumerate(cells):
        set_cell(row[j], v, align=(aligns[j] if aligns else None), size=9)
        if zebra: shade(row[j], CREAM_HEX)
    return row

def fix_widths(t):
    for row in t.rows:
        for j, w in enumerate(t._widths):
            row.cells[j].width = Inches(w)

CTR = WD_ALIGN_PARAGRAPH.CENTER

# ======== Cover ========
p = doc.add_paragraph(); run(p, "National ICT Database of Liberia (NICTD)", size=20, bold=True, color=NAVY)
p = doc.add_paragraph(); run(p, "Data Mining Phase — Field Data Gathering Design & Survey Instruments", size=13, color=TEAL)
p = doc.add_paragraph(); run(p, "ICT Statistics & Policy Unit, Ministry of Posts & Telecommunications. Sampling and staffing computed from the 15-county population frame; see the companion workbook NICTD-Data-Mining-Phase.xlsx.", size=9.5, color=GREY)

# ======== Headline numbers ========
heading("At a glance (all 15 counties)")
t = table(["Metric", "Value"], [3.6, 2.2], header_fill=NAVY_HEX)
for k, v in [
    ("Household interviews (completed / issued)", f"{totals['completed']:,} / {totals['issued']:,}"),
    ("Enumeration areas (urban / rural)", f"{totals['eas']} ({totals['urban_eas']} / {totals['rural_eas']})"),
    ("Field data collectors (enumerators)", f"{totals['enumerators']}"),
    ("Field supervisors", f"{totals['supervisors']}"),
    ("County research leads", f"{totals['research_lead']}"),
    ("Facility / key-informant collectors", f"{totals['kii']}"),
    ("Research cadre (leads + supervisors + KII)", f"{researchers_total}"),
    ("Total field workforce", f"{totals['field_total']}"),
    ("Duration (design to upload)", "approx. 11 weeks"),
]:
    add_row(t, [k, v], aligns=[None, CTR])
fix_widths(t)

# ======== Layers overview ========
heading("The phase in 7 layers")
for n, txt in [
    ("Layer 0 — Source triage", "Of 30 indicators, only 11 need the field: 7 via household survey, 2 via school visits, 2 via a price audit. The other 19 are operator feeds or desk/administrative requests."),
    ("Layer 1 — Sampling design", "Two-stage stratified cluster sample; 20 households per enumeration area; Cochran n0 approx 384 (95%, +/-5%) with finite-population correction, design effect 1.5 and 12% non-response inflation. Completed interviews range 570 (smallest counties) to 1,160 (Montserrado)."),
    ("Layer 2 — Field teams", "Enumerators = completed / (6 per day x 15 days); 1 supervisor per 5 enumerators; 1 county research lead; 2-3 facility/KII collectors. 118 enumerators + 79 research cadre = 197 field staff."),
    ("Layer 3 — Geography & logistics", "Urban and rural EAs per county proportional to urban share; teams based at county capitals; 4x4 access + travel days for the remote south-east and north-west; CAPI tablets with offline GPS-stamped capture."),
    ("Layer 4 — Timeline", "Counties run in parallel: 2 wk design, 2 wk instruments, 1 wk training, 3-day pilot, 3 wk fieldwork, 2 wk cleaning, 1 wk analysis + upload = approx. 11 weeks."),
    ("Layer 5 — Instruments", "Household ICT questionnaire (6 modules), school facility checklist, and tariff price sheet. Full household instrument below."),
    ("Layer 6 — Field to platform", "CAPI capture to weighting to county estimates to the bulk-upload CSV (indicator_code, county, year, value, source) to Content Studio to validation to review queue to live publish."),
]:
    p = doc.add_paragraph(); run(p, n + ". ", bold=True, color=NAVY, size=10); run(p, txt, size=10)

# ======== County sampling table ========
heading("County sampling & field-team plan")
t = table(["County", "Pop.", "Urban", "Completed", "Issued", "EAs", "Enum.", "Sup.", "Lead", "KII", "Team"],
          [1.05, 0.62, 0.55, 0.78, 0.6, 0.5, 0.55, 0.45, 0.45, 0.45, 0.5], header_fill=TEAL_HEX)
for i, r in enumerate(rows):
    add_row(t, [r["county"], f"{r['population']//1000}k", f"{r['urban_share']:.0%}", r["completed"], r["issued"],
                r["eas"], r["enumerators"], r["supervisors"], r["research_lead"], r["kii"], r["field_total"]],
            zebra=(i % 2 == 1),
            aligns=[None] + [CTR]*10)
trow = add_row(t, ["NATIONAL", f"{sum(r['population'] for r in rows)//1000}k", "-", totals["completed"], totals["issued"],
                   totals["eas"], totals["enumerators"], totals["supervisors"], totals["research_lead"], totals["kii"], totals["field_total"]],
               aligns=[None] + [CTR]*10)
for c in trow:
    for rr in c.paragraphs[0].runs: rr.font.bold = True; rr.font.color.rgb = WHITE
    shade(c, NAVY_HEX)
fix_widths(t)

doc.add_page_break()

# ======== THE QUESTIONNAIRE ========
p = doc.add_paragraph(); run(p, "NICTD Household ICT Survey — Field Questionnaire", size=16, bold=True, color=NAVY)
para("CAPI instrument (tablet). Codes in [brackets] are stored values. Every question carries a \"Don't know\" [98] and \"Refused\" [99] option. One questionnaire per sampled household; Module D/E individual questions are asked of one randomly selected adult (18+), with E1 asked of one adult man and one adult woman where present.", size=9.5, color=GREY, italic=True)

def q_module(letter, title):
    heading(f"Module {letter} — {title}", size=12, color=TEAL, space_before=12)

def q(num, text, answers=None, note=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    run(p, f"{num}. ", bold=True, size=10, color=NAVY); run(p, text, size=10)
    if answers:
        pa = doc.add_paragraph(); pa.paragraph_format.left_indent = Inches(0.3); pa.paragraph_format.space_after = Pt(2)
        run(pa, "   ".join(answers), size=9, color=GREY)
    if note:
        pn = doc.add_paragraph(); pn.paragraph_format.left_indent = Inches(0.3)
        run(pn, note, size=8.5, italic=True, color=TEAL)

q_module("A", "Identification & Consent")
q("A1", "County / District / Enumeration Area", ["[auto from sample frame]"])
q("A2", "GPS coordinates", ["[auto-captured]"])
q("A3", "Setting", ["Urban [1]", "Rural [2]"])
q("A4", "Read consent script. Does the respondent agree to participate?", ["Yes [1] -> continue", "No [2] -> end & replace household"],
  note="Consent script: voluntary, anonymous, aggregate statistics only, ~20 minutes, may stop anytime.")
q("A5", "Interviewer ID / Date / Start time", ["[auto]"])

q_module("B", "Household Roster")
q("B1", "How many people usually live in this household?", ["[number]"])
q("B2", "List each member: first name, age (years), sex (M[1]/F[2]), highest education completed.", ["[roster grid, one row per member]"],
  note="Roster drives age/sex weighting and identifies eligible adults (18+) and youth (15-24).")
q("B3", "Randomly select ONE adult (18+) for the individual modules.", ["[CAPI auto-selects via Kish grid]"])

q_module("C", "Household ICT Assets")
q("C1", "Does anyone in this household own a working mobile phone?", ["Yes [1]", "No [2] -> skip to C3"])
q("C2", "Is at least one of those phones a smartphone (can connect to the internet)?", ["Yes [1]", "No [2]"],
  note="Produces: smartphone_ownership (share of adults; combined with roster).")
q("C3", "Does this household have any of the following at home?", ["Radio [a]", "Television [b]", "Computer/laptop/tablet [c]", "Fixed internet connection [d]", "None [e]"])
q("C4", "What is the household's main source of electricity?", ["Grid [1]", "Generator [2]", "Solar [3]", "None [4]"])

q_module("D", "Individual Internet Use  (randomly selected adult)")
q("D1", "Have you personally used the internet in the last 3 months, from any device or location?", ["Yes [1]", "No [2] -> skip to D5", "Don't know [98]"],
  note="Produces: internet_penetration, and — split by A3 setting — urban_connectivity / rural_connectivity.")
q("D2", "Where do you usually access the internet?", ["Home [a]", "Mobile data anywhere [b]", "Work [c]", "School [d]", "Public Wi-Fi [e]", "Cyber cafe [f]"])
q("D3", "How often do you use it?", ["Daily [1]", "At least weekly [2]", "At least monthly [3]", "Less than monthly [4]"])
q("D4", "What do you mainly use it for?", ["Calls/messaging [a]", "Social media [b]", "News/info [c]", "Education [d]", "Work/business [e]", "Government services [f]", "Money/payments [g]"])
q("D5", "If not / rarely: what is the main barrier?", ["Too expensive [1]", "No network [2]", "No device [3]", "No skills [4]", "Not interested [5]", "Safety concerns [6]"])

q_module("E", "Gender & Youth Module")
q("E1", "Ask ONE adult man and ONE adult woman in the household: Do you personally use the internet?", ["Man: Yes[1]/No[2]", "Woman: Yes[1]/No[2]"],
  note="Produces: gender_access_gap (male use rate minus female use rate, in percentage points).")
q("E2", "For each household member aged 15-24: Have you used the internet in the last 3 months?", ["Yes [1]", "No [2]", "per youth member"],
  note="Produces: youth_access (internet use among 15-24 year-olds).")

q_module("F", "Affordability & Income")
q("F1", "In a usual month, about how much does this household spend on mobile airtime and data together?", ["[amount LRD or USD]", "None [0]"])
q("F2", "In a usual month, about how much on any fixed home internet?", ["[amount]", "None [0]"])
q("F3", "Which band best describes this household's total monthly income?", ["Lowest 20% [1]", "2nd [2]", "Middle [3]", "4th [4]", "Highest 20% [5]", "Refused [99]"],
  note="Produces: low_income_access (internet-use rate within band [1]); F1/F2 feed the affordability baskets alongside GNI per capita.")
q("F4", "Record end time. Thank the respondent.", ["[auto]"])

# ======== Facility checklist + price sheet ========
doc.add_page_break()
p = doc.add_paragraph(); run(p, "Companion Instruments", size=16, bold=True, color=NAVY)

heading("School Facility Checklist (Key-Informant)", size=12, color=TEAL)
para("Completed by a KII collector at each sampled school (EMIS frame). Produces school_connectivity and student_device_ratio.", size=9.5, color=GREY, italic=True)
for n, txt, a in [
    ("S1", "School name / EMIS code / county / setting", "[from frame]"),
    ("S2", "Is there a working internet connection at the school today?", "Yes [1]   Sometimes/unreliable [2]   No [3]"),
    ("S3", "Total enrolment (pupils)", "[number]"),
    ("S4", "Number of internet-capable learning devices available to pupils", "[number]"),
    ("S5", "Head-teacher: how is the connection used for teaching/administration?", "[short text]"),
    ("S6", "Main barrier to connectivity", "Cost [1]  No network [2]  No devices [3]  No power [4]  Other [5]"),
]:
    q(n, txt, [a])

heading("Tariff Price Sheet (Price Audit)", size=12, color=TEAL, space_before=10)
para("Completed in each county capital from each operator/ISP shop. Feeds mobile_basket_price and broadband_basket_price (with GNI per capita).", size=9.5, color=GREY, italic=True)
for n, txt, a in [
    ("P1", "Operator / ISP name", "[text]"),
    ("P2", "Cheapest entry-level mobile bundle meeting the standard basket (data + voice), monthly price", "[amount]"),
    ("P3", "Cheapest entry-level fixed-broadband plan, monthly price", "[amount]"),
    ("P4", "Data allowance / speed of the quoted plans", "[GB / Mbps]"),
    ("P5", "Date & location of price capture", "[auto]"),
]:
    q(n, txt, [a])

path = os.path.abspath(os.path.join(OUT_DIR, "NICTD-Data-Mining-Phase.docx"))
doc.save(path)
print("WROTE", path)
